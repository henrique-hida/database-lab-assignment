#!/usr/bin/env python3
"""Gera a carga SQL das tabelas de staging a partir da planilha da Tsusho.

Uso:
    python generate_staging_load.py tsusho-2026-preenchido.xlsx

Também é possível definir o arquivo de saída e o período:
    python generate_staging_load.py tsusho-2026-preenchido.xlsx \
        --saida 07_staging_load_jan_mar.sql \
        --inicio 2026-01-01 \
        --fim 2026-03-31

Dependência:
    pip install openpyxl
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


TipoColuna = Literal["texto", "inteiro", "dinheiro", "data"]


@dataclass(frozen=True)
class Coluna:
    cabecalho_excel: str
    nome_sql: str
    tipo: TipoColuna


@dataclass(frozen=True)
class ConfiguracaoAba:
    aba: str
    tabela: str
    coluna_filtro: str
    colunas: tuple[Coluna, ...]


CONFIGURACOES = (
    ConfiguracaoAba(
        aba="vendas",
        tabela="staging.vendas",
        coluna_filtro="Data do Pedido",
        colunas=(
            Coluna("Cliente", "stv_cliente", "texto"),
            Coluna("CPF ou CNPJ", "stv_cpf_cnpj", "texto"),
            Coluna("Telefone", "stv_telefone", "texto"),
            Coluna("Cidade", "stv_cidade", "texto"),
            Coluna("Modelo", "stv_modelo", "texto"),
            Coluna("Versão", "stv_versao", "texto"),
            Coluna("Cor", "stv_cor", "texto"),
            Coluna("Ano", "stv_ano", "inteiro"),
            Coluna("Placa", "stv_placa", "texto"),
            Coluna("Valor do Carro", "stv_valor_carro", "dinheiro"),
            Coluna("Desconto", "stv_desconto", "dinheiro"),
            Coluna("Valor Final", "stv_valor_final", "dinheiro"),
            Coluna("Forma de Pagamento", "stv_forma_pagamento", "texto"),
            Coluna("Entrada", "stv_entrada", "dinheiro"),
            Coluna("Parcelas", "stv_parcelas", "inteiro"),
            Coluna("Valor Parcela", "stv_valor_parcela", "dinheiro"),
            Coluna("Carro de Troca", "stv_carro_troca", "texto"),
            Coluna("Valor Troca", "stv_valor_troca", "dinheiro"),
            Coluna("Data do Pedido", "stv_data_pedido", "data"),
            Coluna("Data de Faturamento", "stv_data_faturamento", "data"),
            Coluna("Data de Entrega", "stv_data_entrega", "data"),
            Coluna("Status", "stv_status", "texto"),
            Coluna("Observações", "stv_observacoes", "texto"),
        ),
    ),
    ConfiguracaoAba(
        aba="carros",
        tabela="staging.carros",
        coluna_filtro="Entrada no Estoque",
        colunas=(
            Coluna("Modelo", "stc_modelo", "texto"),
            Coluna("Versão", "stc_versao", "texto"),
            Coluna("Ano", "stc_ano", "inteiro"),
            Coluna("Cor", "stc_cor", "texto"),
            Coluna("Placa", "stc_placa", "texto"),
            Coluna("Combustível", "stc_combustivel", "texto"),
            Coluna("Câmbio", "stc_cambio", "texto"),
            Coluna("Valor Tabela", "stc_valor_tabela", "dinheiro"),
            Coluna("Valor Mínimo", "stc_valor_minimo", "dinheiro"),
            Coluna("Entrada no Estoque", "stc_entrada_estoque", "data"),
            Coluna("Data de Reserva", "stc_data_reserva", "data"),
            Coluna("Data de Saída", "stc_data_saida", "data"),
            Coluna("Status", "stc_status", "texto"),
            Coluna("Local", "stc_local", "texto"),
            Coluna("Obs", "stc_observacao", "texto"),
        ),
    ),
    ConfiguracaoAba(
        aba="acessorios",
        tabela="staging.acessorios",
        coluna_filtro="Data do Pedido",
        colunas=(
            Coluna("Cliente", "sta_cliente", "texto"),
            Coluna("Modelo do carro", "sta_modelo_carro", "texto"),
            Coluna("Placa", "sta_placa", "texto"),
            Coluna("Acessório", "sta_acessorio", "texto"),
            Coluna("Categoria", "sta_categoria", "texto"),
            Coluna("Marca", "sta_marca", "texto"),
            Coluna("Quantidade", "sta_quantidade", "inteiro"),
            Coluna("Valor Unitário", "sta_valor_unitario", "dinheiro"),
            Coluna("Desconto", "sta_desconto", "dinheiro"),
            Coluna("Total", "sta_total", "dinheiro"),
            Coluna("Forma de Pagamento", "sta_forma_pagamento", "texto"),
            Coluna("Data do Pedido", "sta_data_pedido", "data"),
            Coluna("Data de instalação", "sta_data_instalacao", "data"),
            Coluna("Status", "sta_status", "texto"),
            Coluna("Observação", "sta_observacao", "texto"),
        ),
    ),
)


def vazio(valor: Any) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def converter_data(valor: Any) -> date | None:
    if vazio(valor):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float, Decimal)):
        convertido = from_excel(float(valor))
        return convertido.date() if isinstance(convertido, datetime) else convertido
    if isinstance(valor, str):
        return date.fromisoformat(valor.strip())
    raise ValueError(f"Data inválida: {valor!r}")


def converter_decimal(valor: Any) -> Decimal:
    try:
        return Decimal(str(valor))
    except (InvalidOperation, ValueError) as erro:
        raise ValueError(f"Número inválido: {valor!r}") from erro


def valor_sql(valor: Any, tipo: TipoColuna) -> str:
    if vazio(valor):
        return "NULL"

    if tipo == "data":
        valor_data = converter_data(valor)
        return f"DATE '{valor_data.isoformat()}'"

    if tipo == "inteiro":
        numero = converter_decimal(valor)
        if numero != numero.to_integral_value():
            raise ValueError(f"Era esperado um inteiro, mas foi recebido {valor!r}")
        return str(int(numero))

    if tipo == "dinheiro":
        return f"{converter_decimal(valor):.2f}"

    texto = str(valor).replace("'", "''")
    return f"'{texto}'"


def ler_e_filtrar_abas(
    caminho_excel: Path,
    data_inicio: date,
    data_fim: date,
) -> dict[str, list[tuple[Any, ...]]]:
    workbook = load_workbook(caminho_excel, read_only=True, data_only=True)
    selecionadas: dict[str, list[tuple[Any, ...]]] = {}

    try:
        for configuracao in CONFIGURACOES:
            if configuracao.aba not in workbook.sheetnames:
                raise ValueError(f"Aba ausente no Excel: {configuracao.aba}")

            worksheet = workbook[configuracao.aba]
            linhas = list(worksheet.iter_rows(values_only=True))
            if not linhas:
                raise ValueError(f"Aba vazia: {configuracao.aba}")

            esperados = tuple(coluna.cabecalho_excel for coluna in configuracao.colunas)
            encontrados = tuple(linhas[0][: len(esperados)])
            if encontrados != esperados:
                raise ValueError(
                    f"Cabeçalhos inesperados na aba {configuracao.aba}.\n"
                    f"Esperados: {esperados}\nEncontrados: {encontrados}"
                )

            indice_filtro = esperados.index(configuracao.coluna_filtro)
            registros = []
            for linha in linhas[1:]:
                linha = tuple(linha[: len(esperados)])
                data_registro = converter_data(linha[indice_filtro])
                if data_registro is not None and data_inicio <= data_registro <= data_fim:
                    registros.append(linha)

            if not registros:
                raise ValueError(f"Nenhum registro selecionado na aba {configuracao.aba}")
            selecionadas[configuracao.aba] = registros
    finally:
        workbook.close()

    return selecionadas


def validar_dados(registros: dict[str, list[tuple[Any, ...]]]) -> None:
    vendas = registros["vendas"]
    carros = registros["carros"]
    acessorios = registros["acessorios"]

    placas_carros = {linha[4] for linha in carros}
    clientes_venda = {linha[8]: linha[0] for linha in vendas}

    vendas_sem_carro = sorted({linha[8] for linha in vendas} - placas_carros)
    if vendas_sem_carro:
        raise ValueError(f"Vendas sem carro correspondente: {', '.join(vendas_sem_carro)}")

    acessorios_sem_venda = sorted(
        {linha[2] for linha in acessorios} - set(clientes_venda)
    )
    if acessorios_sem_venda:
        raise ValueError(
            f"Acessórios sem venda correspondente: {', '.join(acessorios_sem_venda)}"
        )

    for linha in acessorios:
        cliente_venda = clientes_venda[linha[2]]
        if linha[0] != cliente_venda:
            raise ValueError(
                f"Cliente divergente para a placa {linha[2]}: "
                f"venda={cliente_venda!r}, acessório={linha[0]!r}"
            )

    for linha in vendas:
        calculado = converter_decimal(linha[9]) - converter_decimal(linha[10])
        informado = converter_decimal(linha[11])
        if calculado != informado:
            raise ValueError(f"Valor final incorreto na venda da placa {linha[8]}")

    for linha in acessorios:
        calculado = (
            converter_decimal(linha[6]) * converter_decimal(linha[7])
            - converter_decimal(linha[8])
        )
        informado = converter_decimal(linha[9])
        if calculado != informado:
            raise ValueError(
                f"Total incorreto no acessório {linha[3]!r} da placa {linha[2]}"
            )


def montar_insert(
    configuracao: ConfiguracaoAba,
    registros: list[tuple[Any, ...]],
) -> str:
    colunas_sql = ",\n".join(f"    {coluna.nome_sql}" for coluna in configuracao.colunas)
    valores = []

    for linha in registros:
        campos = [
            valor_sql(linha[indice], coluna.tipo)
            for indice, coluna in enumerate(configuracao.colunas)
        ]
        valores.append(f"    ({', '.join(campos)})")

    return (
        f"-- {configuracao.aba}: {len(registros)} registros\n"
        f"INSERT INTO {configuracao.tabela} (\n"
        f"{colunas_sql}\n"
        f") VALUES\n"
        f"{',\n'.join(valores)};"
    )


def gerar_sql(
    caminho_excel: Path,
    caminho_saida: Path,
    data_inicio: date,
    data_fim: date,
) -> dict[str, int]:
    registros = ler_e_filtrar_abas(caminho_excel, data_inicio, data_fim)
    validar_dados(registros)

    secoes = [
        montar_insert(configuracao, registros[configuracao.aba])
        for configuracao in CONFIGURACOES
    ]

    cabecalho = f"""-- ============================================================================
-- PROJETO: Banco de dados da concessionaria Toyota Tsusho
-- ARQUIVO: {caminho_saida.name}
-- DESCRICAO: Carga das tabelas de staging de {data_inicio.isoformat()} a {data_fim.isoformat()}
-- PRE-REQUISITO: 06_staging_tables.sql
-- FONTE: {caminho_excel.name}
-- ============================================================================

BEGIN;
"""
    conteudo = f"{cabecalho}\n{'\n\n'.join(secoes)}\n\nCOMMIT;\n"
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    caminho_saida.write_text(conteudo, encoding="utf-8")

    return {aba: len(linhas) for aba, linhas in registros.items()}


def criar_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Gera os INSERTs das tabelas de staging a partir do Excel."
    )
    parser.add_argument("excel", type=Path, help="Caminho do arquivo .xlsx")
    parser.add_argument(
        "-o",
        "--saida",
        type=Path,
        default=Path("07_staging_load_jan_mar.sql"),
        help="Arquivo SQL de saída",
    )
    parser.add_argument(
        "--inicio",
        type=date.fromisoformat,
        default=date(2026, 1, 1),
        help="Data inicial no formato AAAA-MM-DD",
    )
    parser.add_argument(
        "--fim",
        type=date.fromisoformat,
        default=date(2026, 3, 31),
        help="Data final no formato AAAA-MM-DD",
    )
    return parser


def main() -> None:
    argumentos = criar_parser().parse_args()
    if argumentos.inicio > argumentos.fim:
        raise ValueError("A data inicial não pode ser posterior à data final")
    if argumentos.excel.suffix.lower() != ".xlsx":
        raise ValueError("O arquivo de entrada precisa possuir a extensão .xlsx")

    contagens = gerar_sql(
        argumentos.excel,
        argumentos.saida,
        argumentos.inicio,
        argumentos.fim,
    )

    print(f"Arquivo gerado: {argumentos.saida.resolve()}")
    print(f"Vendas: {contagens['vendas']}")
    print(f"Carros: {contagens['carros']}")
    print(f"Acessórios: {contagens['acessorios']}")


if __name__ == "__main__":
    main()